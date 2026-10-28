from __future__ import annotations

import argparse
from io import BytesIO
import json
import math
import mimetypes
import posixpath
import re
import threading
import webbrowser
import zipfile
from collections import Counter
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
import xml.etree.ElementTree as ET

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(
        "Falta openpyxl. Instálalo con: python -m pip install -r requirements.txt"
    ) from exc


ROOT = Path(__file__).resolve().parent
STATIC = ROOT / "static"
MAIN_SHEET = "RCV"
MAX_REQUEST_BYTES = 16_384
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_PAGE_SIZE = 100
FILTER_COLUMN_IDS = ("C", "D", "E", "CD")
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _resolve_sheet_xml(archive: zipfile.ZipFile, wanted_name: str) -> str:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationships = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in relationships_root.findall(f"{{{PKG_REL_NS}}}Relationship")
    }
    sheets = workbook_root.find(f"{{{MAIN_NS}}}sheets")
    if sheets is None:
        raise ValueError("El libro no tiene una lista de hojas.")

    for sheet in sheets.findall(f"{{{MAIN_NS}}}sheet"):
        if sheet.attrib.get("name") == wanted_name:
            target = relationships.get(sheet.attrib.get(f"{{{REL_NS}}}id", ""))
            if not target:
                raise ValueError(f"No se pudo localizar la hoja {wanted_name}.")
            if target.startswith("/"):
                return target.lstrip("/")
            return posixpath.normpath(posixpath.join("xl", target))

    raise ValueError(f"No existe la hoja {wanted_name}.")


def _last_populated_row(source: Path | bytes, sheet_name: str) -> int:
    """Find the last row with a value or formula, ignoring Excel's styled blank rows."""
    row_tag = f"{{{MAIN_NS}}}row"
    cell_tag = f"{{{MAIN_NS}}}c"
    value_tag = f"{{{MAIN_NS}}}v"
    inline_tag = f"{{{MAIN_NS}}}is"
    formula_tag = f"{{{MAIN_NS}}}f"
    last_row = 0

    archive_source = BytesIO(source) if isinstance(source, bytes) else source
    with zipfile.ZipFile(archive_source) as archive:
        sheet_xml = _resolve_sheet_xml(archive, sheet_name)
        with archive.open(sheet_xml) as stream:
            for _, element in ET.iterparse(stream, events=("end",)):
                if element.tag == row_tag:
                    populated = False
                    for cell in element.findall(cell_tag):
                        value = cell.find(value_tag)
                        if (
                            cell.find(formula_tag) is not None
                            or (value is not None and value.text not in (None, ""))
                            or cell.find(inline_tag) is not None
                        ):
                            populated = True
                            break
                    if populated:
                        last_row = max(last_row, int(element.attrib.get("r", "0")))
                    element.clear()

    return last_row


def _json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _display_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _choose_default_columns(columns: list[dict]) -> list[str]:
    preferred = (
        "DOCUMENTO",
        "EDAD",
        "Var8.Sexo",
        "Var9.Régimen",
        "ZONIFICACIÓN",
        "Var39.Estadio.de.ERC",
        "Var27.Creatinina#1Semestre",
        "Var28.Hemoglobina.Glicosilada#1Semestre",
        "FORMULA-TFG-CCG",
        "FORMULA-TFG-1CKDEPI",
    )
    selected: list[str] = []
    for hint in preferred:
        found = next(
            (
                column["id"]
                for column in columns
                if hint.casefold() in column["header"].casefold()
                and column["id"] not in selected
            ),
            None,
        )
        if found:
            selected.append(found)
    return selected or [column["id"] for column in columns[:8]]


class WorkbookData:
    def __init__(
        self,
        path: Path | None = None,
        *,
        workbook_bytes: bytes | None = None,
        source_name: str | None = None,
    ):
        if workbook_bytes is None:
            if path is None or not path.is_file():
                raise FileNotFoundError(f"No se encontró el archivo: {path}")
            self.path: Path | None = path.resolve()
            self.source_name = self.path.name
            self.last_modified = datetime.fromtimestamp(
                self.path.stat().st_mtime
            ).strftime("%Y-%m-%d %H:%M")
            source: Path | bytes = self.path
        else:
            self.path = None
            self.source_name = Path(source_name or "Libro.xlsx").name[:255]
            self.last_modified = datetime.now().strftime("%Y-%m-%d %H:%M")
            source = workbook_bytes
        self.records: list[dict] = []
        self.columns: list[dict] = []
        self.by_id: dict[str, dict] = {}
        self.sheet_names: list[str] = []
        self._load(source)

    def _load(self, source: Path | bytes) -> None:
        workbook_source = BytesIO(source) if isinstance(source, bytes) else source
        workbook = openpyxl.load_workbook(
            workbook_source, read_only=True, data_only=True, keep_links=False
        )
        try:
            self.sheet_names = list(workbook.sheetnames)
            if MAIN_SHEET not in workbook.sheetnames:
                raise ValueError(f"El libro no contiene la hoja {MAIN_SHEET}.")

            sheet = workbook[MAIN_SHEET]
            populated_last_row = _last_populated_row(source, MAIN_SHEET)
            if populated_last_row < 2:
                raise ValueError(f"La hoja {MAIN_SHEET} no tiene filas de datos.")
            if populated_last_row > 250_000:
                raise ValueError(
                    "La hoja RCV supera el límite local de 250.000 filas con contenido."
                )

            header_values = next(
                sheet.iter_rows(min_row=1, max_row=1, values_only=True)
            )
            raw_headers: list[tuple[str, str]] = []
            for index, raw_value in enumerate(header_values, start=1):
                header = _display_text(raw_value)
                if header:
                    column_id = openpyxl.utils.get_column_letter(index)
                    raw_headers.append((column_id, header))

            counts = Counter(header for _, header in raw_headers)
            self.columns = [
                {
                    "id": column_id,
                    "header": header,
                    "label": (
                        f"{header} ({column_id})"
                        if counts[header] > 1
                        else header
                    ),
                }
                for column_id, header in raw_headers
            ]
            self.by_id = {column["id"]: column for column in self.columns}
            max_col = max(
                openpyxl.utils.column_index_from_string(column["id"])
                for column in self.columns
            )

            for source_row, values in enumerate(
                sheet.iter_rows(
                    min_row=2,
                    max_row=populated_last_row,
                    max_col=max_col,
                    values_only=True,
                ),
                start=2,
            ):
                record = {
                    column["id"]: _json_value(
                        values[
                            openpyxl.utils.column_index_from_string(column["id"]) - 1
                        ]
                    )
                    for column in self.columns
                }
                if not any(value is not None for value in record.values()):
                    continue
                record["_source_row"] = source_row
                record["_search"] = "\u241f".join(
                    str(value) for value in record.values() if value is not None
                ).casefold()
                self.records.append(record)
        finally:
            workbook.close()

        self.default_columns = _choose_default_columns(self.columns)
        self.document_column = next(
            (
                column["id"]
                for column in self.columns
                if column["header"].casefold() == "documento"
            ),
            None,
        )
        self.formula_tfg_columns = [
            column["id"]
            for column in self.columns
            if "formula-tfg" in column["header"].casefold()
        ]
        self.formula_ref_errors = sum(
            record.get(column_id) == "#REF!"
            for record in self.records
            for column_id in self.formula_tfg_columns
        )
        self.missing_documents = sum(
            not _display_text(record.get(self.document_column))
            for record in self.records
        )
        self.filter_options = self._build_filter_options()
        self.basic_distributions = self._build_basic_distributions()

    def _build_basic_distributions(self) -> list[dict]:
        result = []
        fields = (
            ("C", "Sexo"),
            ("D", "Régimen"),
            ("E", "Zonificación"),
        )
        total = len(self.records)
        for column_id, label in fields:
            column = self.by_id.get(column_id)
            if not column:
                continue
            counts = Counter()
            labels = {}
            for record in self.records:
                value = _display_text(record.get(column_id))
                key = ("value", value.casefold()) if value else ("missing", "")
                counts[key] += 1
                labels[key] = value.upper() if value else "Sin dato"
            items = [
                {
                    "label": labels[key],
                    "count": count,
                    "percent": round((count / total) * 100, 1) if total else 0,
                }
                for key, count in sorted(
                    counts.items(),
                    key=lambda item: (-item[1], labels[item[0]].casefold()),
                )
            ]
            result.append(
                {
                    "id": column_id,
                    "label": label,
                    "total": total,
                    "items": items,
                }
            )
        return result

    def _build_filter_options(self) -> list[dict]:
        result = []
        for column_id in FILTER_COLUMN_IDS:
            column = self.by_id.get(column_id)
            if not column:
                continue
            options = sorted(
                {
                    _display_text(record.get(column_id))
                    for record in self.records
                    if _display_text(record.get(column_id))
                },
                key=str.casefold,
            )
            if 1 < len(options) <= 100:
                result.append(
                    {
                        "id": column_id,
                        "label": column["label"],
                        "options": options,
                    }
                )
        return result

    def metadata(self) -> dict:
        return {
            "source_name": self.source_name,
            "sheet": MAIN_SHEET,
            "source_sheets": self.sheet_names,
            "record_count": len(self.records),
            "field_count": len(self.columns),
            "missing_documents": self.missing_documents,
            "formula_ref_errors": self.formula_ref_errors,
            "columns": self.columns,
            "default_columns": self.default_columns,
            "filters": self.filter_options,
            "basic_distributions": self.basic_distributions,
            "last_modified": self.last_modified,
        }

    def query(self, payload: dict) -> dict:
        search = _display_text(payload.get("search"))[:120].casefold()
        filters = payload.get("filters", {})
        if not isinstance(filters, dict):
            filters = {}
        filters = {
            column_id: _display_text(value)
            for column_id, value in filters.items()
            if column_id in self.by_id and value not in (None, "", "__ALL__")
        }

        try:
            page_size = int(payload.get("page_size", 50))
        except (TypeError, ValueError):
            page_size = 50
        page_size = max(10, min(page_size, MAX_PAGE_SIZE))
        try:
            requested_page = int(payload.get("page", 1))
        except (TypeError, ValueError):
            requested_page = 1
        requested_page = max(1, requested_page)

        matches = []
        for record in self.records:
            if search and search not in record["_search"]:
                continue
            if any(
                _display_text(record.get(column_id)) != value
                for column_id, value in filters.items()
            ):
                continue
            matches.append(record)

        total = len(matches)
        pages = max(1, math.ceil(total / page_size))
        page = min(requested_page, pages)
        start = (page - 1) * page_size
        selected_rows = matches[start : start + page_size]
        selected_columns = payload.get("columns", self.default_columns)
        if not isinstance(selected_columns, list):
            selected_columns = self.default_columns
        selected_columns = [
            column_id
            for column_id in selected_columns
            if column_id in self.by_id
        ][:20]
        if not selected_columns:
            selected_columns = self.default_columns[:8]

        return {
            "total": total,
            "page": page,
            "page_size": page_size,
            "rows": [
                {
                    "source_row": record["_source_row"],
                    "values": {
                        column_id: record.get(column_id)
                        for column_id in selected_columns
                    },
                }
                for record in selected_rows
            ],
        }

    def record(self, source_row: int) -> dict | None:
        record = next(
            (
                item
                for item in self.records
                if item["_source_row"] == source_row
            ),
            None,
        )
        if record is None:
            return None
        return {
            "source_row": source_row,
            "values": {
                column["id"]: record.get(column["id"])
                for column in self.columns
            },
        }


class LocalHandler(BaseHTTPRequestHandler):
    data: WorkbookData
    port: int

    def log_message(self, _format, *_args) -> None:
        # Avoid writing request details, including search terms, to the terminal.
        return

    def _host_is_local(self) -> bool:
        host = self.headers.get("Host", "").lower()
        return host in {
            f"127.0.0.1:{self.port}",
            f"localhost:{self.port}",
        }

    def _origin_is_local(self) -> bool:
        origin = self.headers.get("Origin")
        return origin is None or origin in {
            f"http://127.0.0.1:{self.port}",
            f"http://localhost:{self.port}",
        }

    def _send(self, status: int, body: bytes, content_type: str) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header(
            "Content-Security-Policy",
            "default-src 'self'; script-src 'self'; style-src 'self'; "
            "connect-src 'self'; img-src 'self' data:; frame-ancestors 'none'; "
            "base-uri 'none'; form-action 'self'",
        )
        self.end_headers()
        self.wfile.write(body)

    def _send_json(self, status: int, payload: dict) -> None:
        body = json.dumps(
            payload, ensure_ascii=False, allow_nan=False, separators=(",", ":")
        ).encode("utf-8")
        self._send(status, body, "application/json; charset=utf-8")

    def _check_local_request(self) -> bool:
        if not self._host_is_local() or not self._origin_is_local():
            self._send_json(403, {"error": "Acceso local únicamente."})
            return False
        return True

    def do_GET(self) -> None:
        if not self._check_local_request():
            return
        parsed_url = urlsplit(self.path)
        path = parsed_url.path
        if path == "/api/meta":
            self._send_json(200, self.data.metadata())
            return

        files = {
            "/": STATIC / "index.html",
            "/styles.css": STATIC / "styles.css",
            "/app.js": STATIC / "app.js",
        }
        file_path = files.get(path)
        if file_path is None or not file_path.is_file():
            self._send_json(404, {"error": "No encontrado."})
            return
        content_type = mimetypes.guess_type(file_path.name)[0] or "text/plain"
        self._send(200, file_path.read_bytes(), f"{content_type}; charset=utf-8")

    def do_POST(self) -> None:
        if not self._check_local_request():
            return
        parsed_url = urlsplit(self.path)
        path = parsed_url.path
        try:
            content_length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            self._send_json(400, {"error": "Solicitud inválida."})
            return
        if path == "/api/upload":
            self._upload_workbook(content_length, parsed_url.query)
            return
        if content_length < 0 or content_length > MAX_REQUEST_BYTES:
            self._send_json(413, {"error": "La solicitud excede el límite local."})
            return
        try:
            payload = json.loads(self.rfile.read(content_length) or b"{}")
        except (json.JSONDecodeError, UnicodeDecodeError):
            self._send_json(400, {"error": "JSON inválido."})
            return
        if not isinstance(payload, dict):
            self._send_json(400, {"error": "Solicitud inválida."})
            return

        if path == "/api/query":
            self._send_json(200, self.data.query(payload))
            return
        if path == "/api/record":
            try:
                source_row = int(payload.get("source_row", 0))
            except (TypeError, ValueError):
                source_row = 0
            record = self.data.record(source_row)
            if record is None:
                self._send_json(404, {"error": "Registro no encontrado."})
            else:
                self._send_json(200, record)
            return
        self._send_json(404, {"error": "No encontrado."})

    def _upload_workbook(self, content_length: int, query: str) -> None:
        if content_length <= 0:
            self._send_json(400, {"error": "Selecciona un archivo Excel para cargar."})
            return
        if content_length > MAX_UPLOAD_BYTES:
            self._send_json(413, {"error": "El archivo supera el límite local de 50 MB."})
            return

        filename = parse_qs(query).get("filename", ["Libro.xlsx"])[0]
        filename = Path(filename).name[:255]
        if Path(filename).suffix.casefold() not in {".xlsx", ".xlsm"}:
            self._send_json(415, {"error": "Selecciona un libro .xlsx o .xlsm."})
            return

        workbook_bytes = self.rfile.read(content_length)
        if len(workbook_bytes) != content_length:
            self._send_json(400, {"error": "La carga del archivo quedó incompleta. Inténtalo de nuevo."})
            return
        try:
            loaded_data = WorkbookData(
                workbook_bytes=workbook_bytes,
                source_name=filename,
            )
        except Exception as exc:
            message = str(exc).strip() or "El libro no se pudo interpretar."
            self._send_json(400, {"error": f"No se pudo cargar el Excel: {message}"})
            return

        type(self).data = loaded_data
        self._send_json(200, loaded_data.metadata())


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Consulta local y de solo lectura del libro RCV."
    )
    parser.add_argument(
        "--workbook",
        required=True,
        help="Ruta local del archivo .xlsx. El libro no se modifica.",
    )
    parser.add_argument("--port", type=int, default=8765)
    args = parser.parse_args()
    if not 1 <= args.port <= 65535:
        raise SystemExit("El puerto debe estar entre 1 y 65535.")

    try:
        data = WorkbookData(Path(args.workbook).expanduser())
    except Exception as exc:
        raise SystemExit(f"No se pudo abrir el libro: {exc}") from exc

    handler = type(
        "BoundLocalHandler",
        (LocalHandler,),
        {"data": data, "port": args.port},
    )
    server = ThreadingHTTPServer(("127.0.0.1", args.port), handler)
    url = f"http://127.0.0.1:{args.port}/"
    print(f"Libro listo: {data.source_name}")
    print(
        f"Hoja {MAIN_SHEET}: {len(data.records):,} registros, "
        f"{len(data.columns)} campos."
    )
    if data.formula_ref_errors:
        print(
            f"Aviso: {data.formula_ref_errors} resultados #REF! en columnas TFG."
        )
    print(f"Aplicación local: {url} (Ctrl+C para detener)")
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nAplicación detenida.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
